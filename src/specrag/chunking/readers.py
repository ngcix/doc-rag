"""Document text readers."""

from __future__ import annotations

import csv
from pathlib import Path


def read_document(path: Path) -> str:
    suffix = path.suffix.lower()
    readers = {
        ".pdf": _read_pdf,
        ".docx": _read_docx,
        ".xlsx": _read_xlsx,
        ".xls": _read_xls,
        ".pptx": _read_pptx,
        ".csv": _read_csv,
    }
    if suffix in {".txt", ".md", ".markdown"}:
        return path.read_text(encoding="utf-8")
    if reader := readers.get(suffix):
        return reader(path)
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError as exc:
        raise ValueError(f"Unsupported file type: {suffix}") from exc


def _read_pdf(path: Path) -> str:
    try:
        import fitz
    except ImportError as exc:
        raise ImportError("PyMuPDF is required for PDF processing") from exc
    with fitz.open(str(path)) as document:
        return "\n\n".join(text for page in document if (text := page.get_text()))


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise ImportError("python-docx is required for DOCX processing") from exc
    from docx.table import Table

    blocks = []
    for block in Document(str(path)).iter_inner_content():
        text = _docx_table_text(block) if isinstance(block, Table) else block.text
        if text.strip():
            blocks.append(text)
    return "\n\n".join(blocks)


def _docx_table_text(table) -> str:
    lines = []
    for row in table.rows:
        cells = []
        previous = None
        for cell in row.cells:
            # Horizontally merged cells repeat the same underlying cell; keep one.
            if previous is not None and cell._tc is previous._tc:
                continue
            cells.append(cell.text.strip())
            previous = cell
        text = "\t".join(cells)
        if text.strip():
            lines.append(text)
    return "\n".join(lines)


def _read_csv(path: Path) -> str:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        lines = ("\t".join(row) for row in csv.reader(handle))
        return "\n".join(line for line in lines if line.strip())


def _read_xlsx(path: Path) -> str:
    try:
        import openpyxl
    except ImportError as exc:
        raise ImportError("openpyxl is required for XLSX processing") from exc
    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        lines = []
        for name in workbook.sheetnames:
            lines.append(f"=== Sheet: {name} ===")
            for row in workbook[name].iter_rows(values_only=True):
                text = "\t".join("" if cell is None else str(cell) for cell in row)
                if text.strip():
                    lines.append(text)
        return "\n".join(lines)
    finally:
        workbook.close()


def _read_xls(path: Path) -> str:
    try:
        import xlrd
    except ImportError as exc:
        raise ImportError("xlrd is required for XLS processing") from exc
    workbook = xlrd.open_workbook(str(path))
    lines = []
    for sheet in workbook.sheets():
        lines.append(f"=== Sheet: {sheet.name} ===")
        for row in range(sheet.nrows):
            text = "\t".join(str(sheet.cell_value(row, column)) for column in range(sheet.ncols))
            if text.strip():
                lines.append(text)
    return "\n".join(lines)


def _read_pptx(path: Path) -> str:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise ImportError("python-pptx is required for PPTX processing") from exc
    lines = []
    for index, slide in enumerate(Presentation(str(path)).slides, start=1):
        lines.append(f"=== Slide {index} ===")
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines.extend(paragraph.text.strip() for paragraph in shape.text_frame.paragraphs if paragraph.text.strip())
            if shape.has_table:
                lines.extend("\t".join(cell.text for cell in row.cells) for row in shape.table.rows)
    return "\n".join(line for line in lines if line.strip())
