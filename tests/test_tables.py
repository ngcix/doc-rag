"""Spreadsheet tables: merged multi-row headers, record chunks and exact lookup."""

from __future__ import annotations

import openpyxl

from specrag.chunking.chunker import DocumentChunker
from specrag.tables import get_table_store, index_document_tables
from specrag.tables.parse import parse_grid, read_spreadsheet
from specrag.tables.render import sheet_cells, sheet_chunks
from specrag.tables.store import TableStore

# Mirrors the channel-plan layout: label columns merged over the 4 header rows, a band
# title on row 1, sub-band names only in their first column, an empty separator column.
HEADER = [
    ["#", "Country", "", "Alpha-3", "Code", "5GHz PLAN", "", "", "", ""],
    ["", "", "", "", "", "U-NII-1", "", "", "U-NII-2C", ""],
    ["", "", "", "", "", "5180", "5200", "", "5700", "5720"],
    ["", "", "", "", "", "36", "40", "", "140", "144"],
]
MERGES = [(0, 3, 0, 0), (0, 3, 1, 2), (0, 3, 3, 3), (0, 3, 4, 4)]


def _grid() -> list[list[str]]:
    return [
        *HEADER,
        ["1", "그리스", "Greece", "GRC", "GR", "Yes indoor only", "Yes indoor only", "", "Yes", "No"],
        ["2", "영국", "United Kingdom", "GBR", "GB", "Yes", "Yes", "", "Yes", "Yes"],
        ["Note", "1) ISO 3166-1", "", "", "", "", "", "", "", ""],
    ]


def test_merged_multi_row_header_becomes_paths() -> None:
    [table] = parse_grid("5GHz", _grid(), MERGES).tables
    assert table.label_columns == [0, 1, 2, 3, 4]
    assert table.columns[5] == ("5GHz PLAN", "U-NII-1", "5180", "36")
    assert table.columns[9] == ("5GHz PLAN", "U-NII-2C", "5720", "144")
    assert 7 not in table.columns  # separator
    assert [record.row for record in table.records] == [5, 6]


def test_record_chunk_repeats_label_and_compresses_runs() -> None:
    sheet = parse_grid("5GHz", _grid(), MERGES)
    chunks = sheet_chunks("plan.xlsx", [sheet])
    assert chunks[0].content.splitlines() == [
        "plan.xlsx › 5GHz",
        "#: 1 | Country: 그리스 / Greece | Alpha-3: GRC | Code: GR",
        "5GHz PLAN › U-NII-1 › 5180–5200 › 36, 40: Yes indoor only",
        "5GHz PLAN › U-NII-2C › 5700 › 140: Yes",
        "5GHz PLAN › U-NII-2C › 5720 › 144: No",
    ]
    assert chunks[0].metadata == {"sheet": "5GHz", "table": 1, "row": 5, "row_label": "#: 1 | Country: 그리스 / Greece | Alpha-3: GRC | Code: GR"}
    assert chunks[-1].content == "plan.xlsx › 5GHz › notes\nrow 7: Note | 1) ISO 3166-1"


def test_stacked_tables_and_unnamed_label_column() -> None:
    header = [row[:] for row in HEADER]
    header[0][3] = ""  # label column without a heading
    grid = [*header, _grid()[4], [""] * 10, *header, _grid()[5]]
    merges = MERGES + [(top + 6, bottom + 6, left, right) for top, bottom, left, right in MERGES]
    sheet = parse_grid("v2.1", grid, merges)
    assert [[record.row for record in table.records] for table in sheet.tables] == [[5], [11]]
    assert sheet.tables[0].columns[3] == ("column D",)


def test_merged_title_is_not_a_header() -> None:
    grid = [["Channel plan", "Channel plan", "Channel plan"], ["Version", "2.9.1", ""]]
    sheet = parse_grid("Cover", grid, [(0, 0, 0, 2)])
    assert sheet.tables == []
    assert sheet.notes == [(1, "Channel plan"), (2, "Version | 2.9.1")]


def _workbook(path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "5GHz"
    for row in _grid():
        sheet.append([int(value) if value.isdigit() else value or None for value in row])
    for top, bottom, left, right in MERGES:
        sheet.merge_cells(start_row=top + 1, end_row=bottom + 1, start_column=left + 1, end_column=right + 1)
    workbook.save(path)


def test_xlsx_chunks_and_exact_lookup(tmp_path) -> None:
    path = tmp_path / "plan.xlsx"
    _workbook(path)
    chunks = DocumentChunker().chunk_file(path)
    assert chunks[0].chunk_id == f"{path}_0"
    assert chunks[0].metadata["row_label"] == "#: 1 | Country: 그리스 / Greece | Alpha-3: GRC | Code: GR"

    store = TableStore(tmp_path / "tables.sqlite")
    store.replace_document(str(path), sheet_cells(read_spreadsheet(path)))
    result = store.lookup(row="gr", column="144")
    assert (result["row_match"], result["column_match"]) == ("exact", "exact")
    assert [(m["row"], m["column"], m["value"]) for m in result["matches"]] == [(5, "5GHz PLAN › U-NII-2C › 5720 › 144", "No")]
    assert store.lookup(row="Kingdom", column="U-NII-1")["row_match"] == "contains"
    assert {m["value"] for m in store.lookup(row="Kingdom", column="U-NII-1")["matches"]} == {"Yes"}
    [table] = store.sheets()
    assert table["label_columns"] == ["#", "Country", "Alpha-3", "Code"] and table["rows"] == 2

    store.delete_document(str(path))
    assert store.stats() == {"cells": 0}


def test_csv_uses_first_row_as_header(tmp_path, monkeypatch) -> None:
    from specrag.config import settings

    monkeypatch.setattr(settings.vectordb, "faiss_index_dir", tmp_path / "index")
    path = tmp_path / "pm.csv"
    path.write_text("No,Country,Region\n1,Korea,KOR\n2,Japan,JPN\n", encoding="utf-8")
    assert DocumentChunker().chunk_file(path)[1].content == "pm.csv › pm\nNo: 2 | Country: Japan | Region: JPN"
    assert index_document_tables([path], get_table_store()) == 1
    assert get_table_store().lookup(row="Japan", column="Region")["matches"][0]["value"] == "JPN"


def test_footnote_merged_across_label_columns_is_a_note() -> None:
    grid = _grid()
    grid[6][1:5] = ["1) ISO 3166-1"] * 4
    sheet = parse_grid("5GHz", grid, MERGES + [(6, 6, 1, 4)])
    assert [record.row for record in sheet.tables[0].records] == [5, 6]
    assert sheet.notes == [(7, "Note | 1) ISO 3166-1")]
