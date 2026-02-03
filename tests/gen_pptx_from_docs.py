
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

def create_presentation(output_path):
    prs = Presentation()

    # Helpers
    def add_title_slide(title, subtitle):
        slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(slide_layout)
        slide.shapes.title.text = title
        slide.placeholders[1].text = subtitle

    def add_content_slide(title, items):
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)
        slide.shapes.title.text = title
        
        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame
        tf.word_wrap = True
        
        for i, item in enumerate(items):
            if isinstance(item, list): # Level 2
                for subitem in item:
                    p = tf.add_paragraph()
                    p.text = subitem
                    p.level = 1
            else:
                if i == 0:
                    tf.text = item
                else:
                    p = tf.add_paragraph()
                    p.text = item
                    p.level = 0

    # Slide 1: Title
    add_title_slide(
        "SpecRAG System",
        "Hệ thống RAG tối ưu cho tài liệu đặc tả kỹ thuật\nTổng hợp tài liệu chi tiết từ docs/"
    )

    # Slide 2: Mục tiêu & Tổng quan
    add_content_slide(
        "1. Mục tiêu & Tổng quan",
        [
            "SpecRAG được thiết kế chuyên biệt để xử lý các tài liệu kỹ thuật phức tạp (PDF, DOCX, Markdown).",
            "Mục tiêu chính:",
            [
                "Tra cứu chính xác thông tin từ văn bản và sơ đồ.",
                "Hỗ trợ AI Agent (như Cline) hiểu sâu về cấu trúc hệ thống.",
                "Tối ưu hóa khả năng tìm kiếm đa phương thức (Văn bản + Hình ảnh)."
            ],
            "Đặc trưng: Thông minh, Hiệu năng cao, Dễ tích hợp."
        ]
    )

    # Slide 3: Kiến trúc Hệ thống
    add_content_slide(
        "2. Kiến trúc Hệ thống",
        [
            "Hệ thống bao gồm 4 lớp xử lý chính:",
            [
                "Document Chunking: Sử dụng LangChain để chia nhỏ văn bản bảo toàn ngữ cảnh.",
                "Image Extraction: Trích xuất ảnh toàn trang hoặc phát hiện vùng (figures, tables).",
                "Vectorization: Chuyển đổi văn bản/ảnh thành vector 768 chiều (GTE model).",
                "Search Engine: Engine tìm kiếm Hybrid (Semantic + BM25)."
            ],
            "Lưu trữ: FAISS (Facebook AI Similarity Search) cho tốc độ cực nhanh."
        ]
    )

    # Slide 4: Công nghệ Cốt lõi (Stack)
    add_content_slide(
        "3. Công nghệ Cốt lõi",
        [
            "Embedding Model: Alibaba-NLP/gte-multilingual-base (768-dim).",
            "Vector Database: FAISS (IndexFlatIP, Cosine Similarity).",
            "Xử lý PDF: PyMuPDF (fitz) và Docling (RT-DETR) cho layout analysis.",
            "Ngôn ngữ: Python 3.10+, Pydantic v2.",
            "Tìm kiếm: Hybrid Search (RRF fusion) & Multi-Query."
        ]
    )

    # Slide 5: Quy trình Xử lý (Pipeline)
    add_content_slide(
        "4. Quy trình Xử lý Tài liệu",
        [
            "Bước 1 - Input: Người dùng đặt tài liệu vào thư mục data/documents/.",
            "Bước 2 - Processing: Chạy lệnh 'process'.",
            [
                "Tách nhỏ text (Chunking).",
                "Trích xuất ảnh (Extraction) & Tự động gắn thẻ (Auto-tagging).",
                "Vector hóa toàn bộ nội dung."
            ],
            "Bước 3 - Storage: Lưu Index vào data/faiss_index/.",
            "Bước 4 - Search: Sẵn sàng phục vụ query qua CLI hoặc API."
        ]
    )

    # Slide 6: Trích xuất Hình ảnh Thông minh
    add_content_slide(
        "5. Trích xuất Hình ảnh Thông minh",
        [
            "Hệ thống hỗ trợ 2 chế độ chính:",
            [
                "Full-page (-fp): Lưu ảnh toàn bộ trang tài liệu kèm text layer.",
                "Docling (-dl): Sử dụng AI để nhận diện từng Figure, Table, Chart riêng biệt."
            ],
            "Cơ chế Tag Boosting (10x):",
            [
                "Tự động phát hiện caption (Ví dụ: 'Figure 4.1').",
                "Lặp lại tag 10 lần khi vector hóa để tăng trọng số tìm kiếm."
            ],
            "Bộ lọc chất lượng: Loại bỏ ảnh quá nhỏ, ảnh đen hoặc ảnh rỗng."
        ]
    )

    # Slide 7: Công cụ Tìm kiếm (Search Engine)
    add_content_slide(
        "6. Cơ chế Tìm kiếm Đa tầng",
        [
            "Semantic Search: Tìm kiếm theo ý nghĩa (không cần trùng khớp từ khóa).",
            "Hybrid Search: Kết hợp BM25 (từ khóa chính xác) + Vector Similarity.",
            "Multi-Query Search:",
            [
                "Mở rộng câu hỏi của người dùng thành nhiều biến thể.",
                "Chạy song song và tổng hợp kết quả bằng thuật toán RRF (Reciprocal Rank Fusion)."
            ],
            "Hỗ trợ tìm kiếm Ảnh (-pic): Chỉ trả về các sơ đồ, bảng biểu liên quan."
        ]
    )

    # Slide 8: Quản lý Tri thức (Knowledge)
    add_content_slide(
        "7. Quản lý Tri thức & Tự hoàn thiện",
        [
            "Self-Improving: Người dùng có thể bổ sung kiến thức cho tài liệu đã index.",
            "Tính năng Knowledge:",
            [
                "AddKnowledge: Thêm mô tả chi tiết cho một đoạn văn hoặc hình ảnh.",
                "AddTag: Thủ công gắn thêm từ khóa để tối ưu tìm kiếm."
            ],
            "Metadata Management: Lưu trữ thông tin chi tiết (source file, page, context)."
        ]
    )

    # Slide 9: Hướng dẫn CLI (Lệnh cơ bản)
    add_content_slide(
        "8. Hướng dẫn CLI",
        [
            "Khởi tạo: python -m src.main init",
            "Xử lý: python -m src.main process -fp",
            "Tìm kiếm: python -m src.main search \"nội dung cần tìm\"",
            "Tìm kiếm Hybrid: python -m src.main search -hy \"Figure 4.1\"",
            "Tìm kiếm ảnh: python -m src.main search -pic \"sơ đồ kết nối\"",
            "Thống kê: python -m src.main stats"
        ]
    )

    # Slide 10: Tích hợp & Workflow
    add_content_slide(
        "9. Tích hợp AI Agent",
        [
            "SpecRAG được tích hợp sẵn với Cline Agent (qua .clinerules).",
            "RAG Backend Workflow:",
            [
                "Tự động phân tích câu hỏi phức tạp.",
                "Kết hợp tra cứu văn bản và sơ đồ đồng thời.",
                "Tự động phân tích nội dung ảnh trả về kết quả cuối cùng."
            ],
            "Tiềm năng: Thay thế việc tra cứu thủ công các bộ spec hàng nghìn trang."
        ]
    )

    # Save
    prs.save(output_path)
    print(f"Presentation saved to {output_path}")

if __name__ == "__main__":
    output_dir = "reports"
    os.makedirs(output_dir, exist_ok=True)
    create_presentation(os.path.join(output_dir, "SpecRAG_Detailed_Docs.pptx"))
