"""Document reader tests."""

from __future__ import annotations

from docx import Document

from specrag.chunking.readers import read_document
from specrag.ingestion.loaders import discover_documents


def test_docx_keeps_tables_in_document_order(tmp_path) -> None:
    document = Document()
    document.add_paragraph("Before table")
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged"
    table.cell(0, 2).text = "Header"
    for column, value in enumerate(["a", "b", "c"]):
        table.cell(1, column).text = value
    document.add_paragraph("After table")
    path = tmp_path / "spec.docx"
    document.save(str(path))

    assert read_document(path) == "Before table\n\nMerged\tHeader\na\tb\tc\n\nAfter table"


def test_csv_rows_are_tab_separated(tmp_path) -> None:
    path = tmp_path / "params.csv"
    path.write_bytes('﻿name,value\n"timer, T3412",54\n\n'.encode())

    assert read_document(path) == "name\tvalue\ntimer, T3412\t54"


def test_csv_is_discovered(tmp_path) -> None:
    (tmp_path / "params.csv").write_text("a,b\n", encoding="utf-8")

    assert discover_documents(tmp_path) == [tmp_path / "params.csv"]
