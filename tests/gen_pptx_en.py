
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

def create_presentation_en(output_path):
    prs = Presentation()

    # Helpers
    def add_title_slide(title, subtitle):
        slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(slide_layout)
        slide.shapes.title.text = title
        slide.placeholders[1].text = subtitle

    def add_content_slide(title, items):
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)
        slide.shapes.title.text = title
        
        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame
        tf.word_wrap = True
        
        for i, item in enumerate(items):
            if isinstance(item, list): # Level 2
                for subitem in item:
                    p = tf.add_paragraph()
                    p.text = subitem
                    p.level = 1
            else:
                if i == 0:
                    tf.text = item
                else:
                    p = tf.add_paragraph()
                    p.text = item
                    p.level = 0

    # Slide 1: Title
    add_title_slide(
        "SpecRAG System",
        "Optimized Retrieval-Augmented Generation for Technical Specifications\nDetailed Technical Summary from Documentation"
    )

    # Slide 2: Objectives & Overview
    add_content_slide(
        "1. Objectives & Overview",
        [
            "SpecRAG is specifically designed to handle complex technical documents (PDF, DOCX, Markdown).",
            "Key Objectives:",
            [
                "Accurate retrieval from both text and diagrams.",
                "Empower AI Agents (like Cline) with deep system structure understanding.",
                "Optimize multimodal search (Text + Images)."
            ],
            "Key Characteristics: Intelligent, High-performance, Seamless Integration."
        ]
    )

    # Slide 3: System Architecture
    add_content_slide(
        "2. System Architecture",
        [
            "The system consists of 4 main processing layers:",
            [
                "Document Chunking: Context-aware splitting using LangChain.",
                "Image Extraction: Full-page extraction or ROI detection (figures, tables).",
                "Vectorization: Conversion to 768-dimensional vectors using GTE models.",
                "Search Engine: Hybrid search engine (Semantic + BM25)."
            ],
            "Storage: Powered by FAISS (Facebook AI Similarity Search) for extreme speed."
        ]
    )

    # Slide 4: Core Technology Stack
    add_content_slide(
        "3. Core Technology Stack",
        [
            "Embedding Model: Alibaba-NLP/gte-multilingual-base (768-dim).",
            "Vector Database: FAISS (IndexFlatIP, Cosine Similarity).",
            "PDF Processing: PyMuPDF (fitz) and Docling (RT-DETR) for layout analysis.",
            "Backend: Python 3.10+, Pydantic v2.",
            "Search Algorithms: Hybrid Search (RRF fusion) & Multi-Query."
        ]
    )

    # Slide 5: Document Processing Pipeline
    add_content_slide(
        "4. Document Processing Pipeline",
        [
            "Step 1 - Input: Documents are placed in 'data/documents/'.",
            "Step 2 - Processing: Execution of the 'process' command.",
            [
                "Text Chunking.",
                "Image Extraction & Auto-tagging.",
                "Content Vectorization."
            ],
            "Step 3 - Storage: Indices are saved in 'data/faiss_index/'.",
            "Step 4 - Search: Ready for CLI or API queries."
        ]
    )

    # Slide 6: Intelligent Image Extraction
    add_content_slide(
        "5. Intelligent Image Extraction",
        [
            "Supports two main modes:",
            [
                "Full-page (-fp): Captures entire pages with text layer context.",
                "Docling (-dl): AI-driven detection of Figures, Tables, and Charts."
            ],
            "Tag Boosting Mechanism (10x):",
            [
                "Automatic caption detection (e.g., 'Figure 4.1').",
                "Tags are repeated 10x during vectorization to increase search weight."
            ],
            "Quality Filtering: Filters out tiny, dark, or empty images."
        ]
    )

    # Slide 7: Multi-layered Search Engine
    add_content_slide(
        "6. Multi-layered Search Engine",
        [
            "Semantic Search: Finds concepts even without exact keyword matches.",
            "Hybrid Search: Combines BM25 (keyword accuracy) + Vector Similarity.",
            "Multi-Query Search:",
            [
                "Expands user questions into multiple variations.",
                "Parallel execution merged via RRF (Reciprocal Rank Fusion)."
            ],
            "Image Search Support (-pic): Specifically retrieves relevant diagrams and tables."
        ]
    )

    # Slide 8: Knowledge Management
    add_content_slide(
        "7. Knowledge Management & Self-Improvement",
        [
            "Self-Improving: Users can append knowledge to indexed documents.",
            "Knowledge Features:",
            [
                "AddKnowledge: Append detailed descriptions to text or images.",
                "AddTag: Manually add keywords to optimize retrieval."
            ],
            "Metadata Management: Stores source file, page numbers, and surrounding context."
        ]
    )

    # Slide 9: CLI Reference (Key Commands)
    add_content_slide(
        "8. CLI Reference",
        [
            "Initialize: python -m src.main init",
            "Process: python -m src.main process -fp",
            "Search: python -m src.main search \"query text\"",
            "Hybrid Search: python -m src.main search -hy \"Figure 4.1\"",
            "Image Search: python -m src.main search -pic \"connection diagram\"",
            "Statistics: python -m src.main stats"
        ]
    )

    # Slide 10: Integration & Workflow
    add_content_slide(
        "9. AI Agent Integration",
        [
            "SpecRAG is pre-integrated with Cline Agent (via .clinerules).",
            "RAG Backend Workflow:",
            [
                "Automated complex query analysis.",
                "Simultaneous text and diagram retrieval.",
                "Automated image content analysis for final results."
            ],
            "Potential: Replaces manual lookup of thousand-page specifications."
        ]
    )

    # Save
    prs.save(output_path)
    print(f"English Presentation saved to {output_path}")

if __name__ == "__main__":
    output_dir = "reports"
    os.makedirs(output_dir, exist_ok=True)
    create_presentation_en(os.path.join(output_dir, "SpecRAG_Detailed_Docs_EN.pptx"))
